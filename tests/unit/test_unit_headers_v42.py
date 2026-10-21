"""v4.2 T2·T3: 헤더 형식(작성자·최종 수정자 포함)과 재색인 시 단위 ID·qrels 결정성

- T2: 이메일·채팅·문서·첨부 헤더 형식, 빈 필드 생략, '외 n명/n건' 규칙 (지시서 v4.2 §2 U1, §9 T2)
- T3: 같은 원본을 두 번 색인하면 단위 ID 집합이 같고 기존 qrels.txt가 그대로 적용된다 (§9 T3)

모든 자료는 테스트 안에서 만든 가명 자료다 (P0).
"""

from __future__ import annotations

import io
import sys
import tempfile
from email.message import EmailMessage as PyEmail
from pathlib import Path

import pytest

from src.parsers.document_parser import DocumentParser
from src.units.builder import UnitBuilder
from src.units.header import attachment_header, chat_header, document_header, email_header
from src.units.tokens import TokenCounter

COUNTER = TokenCounter(hf_name=None, chars_per_token=1.0)
XLSX_MIME = ("application", "vnd.openxmlformats-officedocument.spreadsheetml.sheet")


class _NoOcr:
    name = "none"

    def is_available(self) -> bool:
        return False


@pytest.fixture
def builder(monkeypatch: pytest.MonkeyPatch) -> UnitBuilder:
    parser = DocumentParser(ocr_languages="eng")
    monkeypatch.setattr(parser, "_get_ocr_engine", lambda: _NoOcr())
    return UnitBuilder(
        "case1",
        counter=COUNTER,
        max_part_tokens=500,
        doc_parser=parser,
        clean_enabled=True,
        gap_minutes=30,
    )


def _docx_bytes(author: str, modifier: str) -> bytes:
    from docx import Document

    d = Document()
    d.add_paragraph("정산 기준 조항: 선급금은 검수 후 지급한다.")
    d.core_properties.author = author
    d.core_properties.last_modified_by = modifier
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _xlsx_bytes(creator: str, modifier: str) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "견적"
    ws.append(["품목", "금액"])
    ws.append(["부품A", 1000])
    wb.properties.creator = creator
    wb.properties.lastModifiedBy = modifier
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pdf_bytes(author: str) -> bytes:
    import fitz

    doc = fitz.open()
    doc.new_page().insert_text((72, 72), "payment approval text " * 3)
    doc.set_metadata({"author": author})
    data = doc.tobytes()
    doc.close()
    return data


def _eml_bytes(
    attachments: list[tuple[str, bytes, tuple[str, str]]] | None = None,
    message_id: str = "<H1@x.com>",
) -> bytes:
    m = PyEmail()
    m["From"] = "홍길동 <gd@x.com>"
    m["To"] = "김철수 <cs@x.com>"
    m["Subject"] = "RE: 3월 정산 건"
    m["Date"] = "Tue, 05 Mar 2024 05:22:00 +0000"
    m["Message-ID"] = message_id
    m.set_content("견적서 확인 부탁드립니다.")
    for name, data, (maintype, subtype) in attachments or []:
        m.add_attachment(data, maintype=maintype, subtype=subtype, filename=name)
    return m.as_bytes()


def _write(root: Path, rel: str, data: bytes) -> Path:
    p = root / rel
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_bytes(data)
    return p


# ─────────────── T2: 헤더 함수 (고정 자료) ───────────────


def test_document_header_full_and_omits_empty():
    meta = {
        "filename": "정산기준_v3.docx",
        "page_kind": "page",
        "page": 3,
        "page_count": 12,
        "author": "홍길동",
        "last_modified_by": "박민수",
        "last_modified": "2024-02-28T10:00:00",
    }
    assert document_header(meta) == (
        "[문서] 정산기준_v3.docx | 3/12쪽 | 작성자: 홍길동 | 최종 수정자: 박민수 | 수정일: 2024-02-28"
    )
    # 최종 수정자 없음(PDF) → 필드 생략, 구분자도 남지 않는다
    meta.pop("last_modified_by")
    assert (
        document_header(meta)
        == "[문서] 정산기준_v3.docx | 3/12쪽 | 작성자: 홍길동 | 수정일: 2024-02-28"
    )
    assert document_header({"filename": "a.txt"}) == "[문서] a.txt"


def test_attachment_header_order_with_parent():
    meta = {
        "filename": "견적서.xlsx",
        "page_kind": "page",
        "page": 2,
        "page_count": 3,
        "author": "김철수",
        "last_modified_by": "김철수",
        "parent_unit_id": "em:0",
        "parent_date": "2024-03-05T05:22:00+00:00",
        "parent_sender": "홍길동",
        "parent_recipients": ["김철수", "이영희"],
        "parent_subject": "RE: 3월 정산 건",
    }
    assert attachment_header(meta) == (
        "[첨부] 견적서.xlsx | 2/3쪽 | 작성자: 김철수 | 최종 수정자: 김철수 | "
        '부모 메일: 2024-03-05 홍길동→김철수 외 1명 "RE: 3월 정산 건"'
    )
    # 부모 없는 첨부·빈 작성자 → 해당 필드 모두 생략
    assert attachment_header({"filename": "x.pdf", "author": ""}) == "[첨부] x.pdf"


def test_email_and_chat_limits():
    names = [f"사람{i}" for i in range(7)]
    h = email_header(
        {
            "date": "2024-03-05T05:22:00+00:00",
            "sender_name": "홍길동",
            "recipient_names": names,
            "cc": [],
            "subject": "",
            "attachment_filenames": [f"f{i}.pdf" for i in range(6)],
        }
    )
    assert h == (
        "[이메일] 2024-03-05 14:22 | 보낸이: 홍길동 | "
        "받는이: 사람0, 사람1, 사람2, 사람3, 사람4 외 2명 | "
        "첨부: f0.pdf, f1.pdf, f2.pdf, f3.pdf, f4.pdf 외 1건"
    )
    c = chat_header(
        {
            "date": "2024-03-05T05:00:00+00:00",
            "date_end": "2024-03-05T05:47:00+00:00",
            "participants": names[:6],
            "room": "",
        }
    )
    assert c == "[채팅] 2024-03-05 14:00~14:47 | 참여자: 사람0, 사람1, 사람2, 사람3, 사람4 외 1명"


# ─────────────── T2: 실제 파서 경로 (core properties → 헤더) ───────────────


def test_docx_core_properties_in_header(builder: UnitBuilder, tmp_path: Path):
    p = _write(tmp_path, "docs/r.docx", _docx_bytes("작성자A", "수정자B"))
    u = builder.build_file(p, "docs/r.docx")[0]
    assert "| 작성자: 작성자A | 최종 수정자: 수정자B" in u.header
    assert u.meta["author"] == "작성자A" and u.meta["last_modified_by"] == "수정자B"


def test_pdf_author_only(builder: UnitBuilder, tmp_path: Path):
    p = _write(tmp_path, "s.pdf", _pdf_bytes("PDF작성자"))
    u = builder.build_file(p, "s.pdf")[0]
    assert "| 작성자: PDF작성자" in u.header
    assert "최종 수정자" not in u.header


def test_xlsx_attachment_header(builder: UnitBuilder, tmp_path: Path):
    eml = _eml_bytes([("견적서.xlsx", _xlsx_bytes("김철수", "이영희"), XLSX_MIME)])
    units = builder.build_file(_write(tmp_path, "m.eml", eml), "m.eml")
    att = [u for u in units if u.unit_type == "attachment"][0]
    assert att.header.startswith(
        "[첨부] 견적서.xlsx | 시트 견적 2~2행 | 작성자: 김철수 | 최종 수정자: 이영희 | 부모 메일: "
    )


# ─────────────── T3: 재색인 결정성 + qrels 적용 ───────────────


@pytest.fixture
def fake_env(monkeypatch: pytest.MonkeyPatch):
    """가짜 임베더 + 임시 색인·DB"""
    from src.db.database import reset_globals
    from src.index import case_index
    from src.indexing import pipeline as pipeline_mod
    from src.utils.config import settings

    with tempfile.TemporaryDirectory() as d:
        root = Path(d)
        monkeypatch.setattr(settings, "ast_fake_models", True)
        monkeypatch.setattr(settings, "embed_dim", 32)
        monkeypatch.setattr(settings, "index_dir", str(root / "index"))
        monkeypatch.setattr(settings, "service_index_profile", "hdr-qwen3")
        monkeypatch.setattr(settings, "ollama_embed_model", "qwen3-embedding:4b-q8_0")
        monkeypatch.setattr(pipeline_mod, "_save_failed_parts", lambda *a, **k: None)
        case_index._cache.clear()
        reset_globals()
        from src.cases.case_store import CaseStore

        yield root, CaseStore(db_url=f"sqlite:///{root / 'test.db'}")
        reset_globals()
        case_index._cache.clear()


def _corpus(root: Path) -> Path:
    src = root / "src"
    _write(
        src,
        "kim/Inbox/a.eml",
        _eml_bytes(
            [
                ("견적서.xlsx", _xlsx_bytes("김철수", "이영희"), XLSX_MIME),
                ("규정.pdf", _pdf_bytes("작성자P"), ("application", "pdf")),
            ]
        ),
    )
    _write(src, "kim/Inbox/b.eml", _eml_bytes(message_id="<H2@x.com>"))
    _write(src, "docs/r.docx", _docx_bytes("작성자A", "수정자B"))
    _write(src, "docs/n.txt", "메모 본문 내용 ".encode() * 30)
    return src


def test_reindex_same_unit_ids_and_qrels_apply(fake_env):
    from src.index import case_index
    from src.indexing.pipeline import IndexingPhase, IndexingPipeline

    root, store = fake_env
    src = _corpus(root)
    pl = IndexingPipeline(case_store=store)

    def build(name: str) -> tuple[str, set[str]]:
        meta = store.create(name, doc_paths=[str(src)])
        assert pl.run(meta.case_id).phase == IndexingPhase.COMPLETED
        return meta.case_id, {p.unit_id for p in case_index.load(meta.case_id).parts}

    cid1, ids1 = build("첫 색인")
    _, ids2 = build("같은 원본 새 케이스")
    assert len(ids1) >= 5 and ids1 == ids2
    assert {u.split(":")[0] for u in ids1} == {"em", "att", "doc"}

    # 같은 케이스를 다시 색인해도 같다
    assert pl.run(cid1).phase == IndexingPhase.COMPLETED
    case_index._cache.clear()
    assert {p.unit_id for p in case_index.load(cid1).parts} == ids1

    # 첫 색인으로 만든 qrels.txt가 두 번째 색인에 그대로 적용된다 (모든 unit_id가 존재)
    sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "experiments_v4"))
    try:
        from ev4.io import read_qrels, write_qrels
    finally:
        sys.path.pop(0)
    qrels = {"q001": {u: (2 if i == 0 else 1) for i, u in enumerate(sorted(ids1))}}
    path = root / "qrels.txt"
    write_qrels(path, qrels)
    loaded = read_qrels(path)
    assert loaded == qrels
    assert set(loaded["q001"]) <= ids2
